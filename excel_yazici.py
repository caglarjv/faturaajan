"""Muhasebe şablonlu Excel dosyasını oluşturur ve kayıt ekler.
Tek başına çalıştırırsanız (python excel_yazici.py) boş şablonu oluşturur.
Eski kayıtlardaki arşiv yollarını tıklanabilir yapmak için: python excel_yazici.py --linkle"""
import datetime
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import BASE, EXCEL_YOLU

# (anahtar, başlık, genişlik, sayı biçimi)
SUTUNLAR = [
    ("kayit_no", "Kayıt No", 9, "0"),
    ("kayit_tarihi", "Kayıt Tarihi", 12, "DD.MM.YYYY"),
    ("fatura_tarihi", "Fatura Tarihi", 12, "DD.MM.YYYY"),
    ("donem", "Dönem", 9, "@"),
    ("yon", "Yön", 8, "@"),
    ("fatura_no", "Fatura No", 20, "@"),
    ("ettn", "ETTN", 20, "@"),
    ("satici_unvan", "Satıcı Unvan", 32, "@"),
    ("satici_vkn", "Satıcı VKN/TCKN", 15, "@"),
    ("satici_vd", "Satıcı Vergi Dairesi", 20, "@"),
    ("alici_unvan", "Alıcı Unvan", 32, "@"),
    ("alici_vkn", "Alıcı VKN/TCKN", 15, "@"),
    ("alici_vd", "Alıcı Vergi Dairesi", 20, "@"),
    ("para_birimi", "Para Birimi", 9, "@"),
    ("matrah", "Matrah (KDV Hariç)", 16, "#,##0.00"),
    ("kdv_orani", "KDV Oranı %", 10, "0"),
    ("kdv_tutari", "KDV Tutarı", 14, "#,##0.00"),
    ("genel_toplam", "Genel Toplam", 16, "#,##0.00"),
    ("odeme", "Ödeme / IBAN", 28, "@"),
    ("hesap_kodu", "Hesap Kodu (elle)", 14, "@"),
    ("kategori", "Gider/Gelir Kategorisi (elle)", 22, "@"),
    ("kaynak", "Kaynak", 10, "@"),
    ("gonderen", "Gönderen", 28, "@"),
    ("dosya_adi", "Dosya Adı", 30, "@"),
    ("arsiv_yolu", "Arşiv Yolu", 40, "@"),
    ("durum", "Durum", 16, "@"),
    ("not", "Kontrol Notu", 50, "@"),
    ("hash", "Dosya Hash", 12, "@"),
]
IDX = {k: i + 1 for i, (k, *_r) in enumerate(SUTUNLAR)}
SON_SUTUN = get_column_letter(len(SUTUNLAR))
SINIR = 20000


class ExcelKilitli(Exception):
    """Excel dosyası açık olduğu için yazılamadı."""


def L(anahtar):
    return get_column_letter(IDX[anahtar])


def sablon_olustur(yol=EXCEL_YOLU):
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    baslik_dolgu = PatternFill("solid", fgColor="1F4E78")
    for i, (_k, baslik, gen, _f) in enumerate(SUTUNLAR, start=1):
        c = ws.cell(row=1, column=i, value=baslik)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = baslik_dolgu
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = gen
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{SON_SUTUN}1"
    ws.column_dimensions[L("hash")].hidden = True

    dv_yon = DataValidation(type="list", formula1='"Alış,Satış,Belirsiz"', allow_blank=True)
    dv_durum = DataValidation(type="list", formula1='"Tamam,Kontrol Gerekli,Düzeltildi"',
                              allow_blank=True)
    ws.add_data_validation(dv_yon)
    ws.add_data_validation(dv_durum)
    dv_yon.add(f"{L('yon')}2:{L('yon')}{SINIR}")
    dv_durum.add(f"{L('durum')}2:{L('durum')}{SINIR}")
    ws.conditional_formatting.add(
        f"A2:{SON_SUTUN}{SINIR}",
        FormulaRule(formula=[f'${L("durum")}2="Kontrol Gerekli"'],
                    fill=PatternFill("solid", bgColor="FFF2CC")),
    )

    # ---- Özet sayfası ----
    oz = wb.create_sheet("Özet")
    oz["A1"], oz["B1"] = "Yıl", datetime.date.today().year
    oz["A2"], oz["B2"] = "Para Birimi", "TRY"
    for h in ("A1", "A2"):
        oz[h].font = Font(bold=True)
    F = "Faturalar!"
    yil = f'{F}${L("donem")}:${L("donem")},$B$1&"-*"'
    pb = f'{F}${L("para_birimi")}:${L("para_birimi")},$B$2'

    def top(sutun, yon, ek=""):
        return (f'SUMIFS({F}${L(sutun)}:${L(sutun)},{F}${L("yon")}:${L("yon")},"{yon}",'
                f'{yil},{pb}{ek})')

    for i, b in enumerate(["Yön", "Fatura Adedi", "Matrah", "KDV", "Genel Toplam"], start=1):
        c = oz.cell(row=4, column=i, value=b)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), baslik_dolgu
    for r, yon in ((5, "Alış"), (6, "Satış")):
        oz.cell(row=r, column=1, value=yon)
        oz.cell(row=r, column=2,
                value=f'=COUNTIFS({F}${L("yon")}:${L("yon")},A{r},{yil},{pb})')
        oz.cell(row=r, column=3, value="=" + top("matrah", yon))
        oz.cell(row=r, column=4, value="=" + top("kdv_tutari", yon))
        oz.cell(row=r, column=5, value="=" + top("genel_toplam", yon))
    oz["A7"], oz["D7"] = "Ödenecek (+) / Devreden (-) KDV", "=D6-D5"
    oz["A7"].font = oz["D7"].font = Font(bold=True)
    oz["A8"] = "Kontrol gerekli kayıt sayısı"
    oz["D8"] = f'=COUNTIF({F}${L("durum")}:${L("durum")},"Kontrol Gerekli")'

    for i, b in enumerate(["Dönem", "Alış Matrah", "Alış KDV", "Satış Matrah",
                           "Satış KDV", "Ödenecek KDV"], start=1):
        c = oz.cell(row=10, column=i, value=b)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), baslik_dolgu
    for n in range(1, 13):
        r = 10 + n
        oz.cell(row=r, column=1, value=f'=$B$1&"-"&TEXT({n},"00")')
        donem = f'{F}${L("donem")}:${L("donem")},$A{r}'

        def aylik(sutun, yon):
            return (f'=SUMIFS({F}${L(sutun)}:${L(sutun)},{F}${L("yon")}:${L("yon")},"{yon}",'
                    f'{donem},{pb})')
        oz.cell(row=r, column=2, value=aylik("matrah", "Alış"))
        oz.cell(row=r, column=3, value=aylik("kdv_tutari", "Alış"))
        oz.cell(row=r, column=4, value=aylik("matrah", "Satış"))
        oz.cell(row=r, column=5, value=aylik("kdv_tutari", "Satış"))
        oz.cell(row=r, column=6, value=f"=E{r}-C{r}")
    for satir in oz.iter_rows(min_row=5, max_row=22, min_col=2, max_col=6):
        for c in satir:
            c.number_format = "#,##0.00"
    for r in (5, 6):
        oz.cell(row=r, column=2).number_format = "0"
    oz["D8"].number_format = "0"
    for col, gen in zip("ABCDEF", (34, 16, 16, 16, 16, 16)):
        oz.column_dimensions[col].width = gen

    # ---- Açıklama sayfası ----
    ac = wb.create_sheet("Açıklama")
    satirlar = [
        "FATURA KAYITLARI - KULLANIM",
        "",
        "Faturalar sayfası: ajan her faturayı otomatik olarak yeni satıra ekler.",
        "Sarı satırlar 'Kontrol Gerekli' demektir: Kontrol Notu sütununu okuyup düzeltin,",
        "sonra Durum sütununu 'Düzeltildi' yapın.",
        "Hesap Kodu ve Kategori sütunlarını muhasebe için elle doldurun (ajan dokunmaz).",
        "Yön: KENDI_VKN .env'de tanımlıysa Alış/Satış otomatik belirlenir.",
        "Özet sayfası: B1'deki yıl ve B2'deki para birimine göre otomatik toplar.",
        "UYARI: Ajan çalışırken Excel dosyası açıksa yazamaz; dosyayı kapatınca bekleyen",
        "faturalar bir sonraki turda otomatik işlenir.",
        "Dosya Hash sütunu (gizli) aynı dosyanın iki kez eklenmesini önler.",
    ]
    for i, s in enumerate(satirlar, start=1):
        ac.cell(row=i, column=1, value=s)
    ac["A1"].font = Font(bold=True, size=13)
    ac.column_dimensions["A"].width = 100

    wb.save(yol)


def _ac():
    if not EXCEL_YOLU.exists():
        sablon_olustur()
    return load_workbook(EXCEL_YOLU)


def mukerrer_mi(hash_, fatura_no=None, satici_vkn=None):
    """Aynı dosya veya aynı (fatura no + satıcı VKN) varsa nedenini döndürür."""
    if not EXCEL_YOLU.exists():
        return None
    try:
        wb = load_workbook(EXCEL_YOLU, read_only=True, data_only=True)
    except PermissionError as e:
        raise ExcelKilitli from e
    try:
        ws = wb["Faturalar"]
        ih, ino, ivkn = IDX["hash"] - 1, IDX["fatura_no"] - 1, IDX["satici_vkn"] - 1
        for satir in ws.iter_rows(min_row=2, values_only=True):
            if len(satir) <= ih:
                continue
            if hash_ and satir[ih] == hash_:
                return "aynı dosya daha önce işlenmiş"
            if fatura_no and satici_vkn and satir[ino] == fatura_no and satir[ivkn] == satici_vkn:
                return f"aynı fatura no ({fatura_no}) ve satıcı zaten kayıtlı"
    finally:
        wb.close()
    return None


def _hedef(yol_metni) -> Path:
    """Arşiv dosyasının güncel yolunu bulur (klasör taşınmışsa 'islenen' kısmından yeniden kurar)."""
    p = Path(str(yol_metni))
    if p.exists():
        return p
    parcalar = p.parts
    if "islenen" in parcalar:
        q = BASE / Path(*parcalar[parcalar.index("islenen"):])
        if q.exists():
            return q
    return p


def _yol_metni(deger) -> str:
    """Hücre düz yol ya da =HYPERLINK("yol","yol") formülü olabilir; yolu döndürür."""
    s = str(deger)
    m = re.match(r'=HYPERLINK\("(.*?)"', s)
    return m.group(1) if m else s


def _linkle(hucre):
    """Hücredeki dosya yolunu, tıklanınca dosyayı açan HYPERLINK formülüne çevirir.
    (file:/// bağlantıları Türkçe karakterli adlarda Excel'de bozuluyor; formül sorunsuz.)"""
    if not hucre.value:
        return
    yol = str(_hedef(_yol_metni(hucre.value)).absolute())
    hucre.hyperlink = None  # önceki sürümden kalan bağlantıyı temizle
    if len(yol) > 250:  # Excel formülündeki metin sınırı 255 karakter
        hucre.value = yol
        return
    hucre.value = f'=HYPERLINK("{yol}","{yol}")'
    hucre.style = "Hyperlink"
    hucre.number_format = "General"


def eski_kayitlari_linkle() -> int:
    wb = _ac()
    ws = wb["Faturalar"]
    sutun = IDX["arsiv_yolu"]
    adet = 0
    for satir in range(2, ws.max_row + 1):
        h = ws.cell(row=satir, column=sutun)
        if h.value:
            _linkle(h)
            adet += 1
    wb.save(EXCEL_YOLU)
    return adet


def ekle(kayit: dict) -> int:
    try:
        wb = _ac()
    except PermissionError as e:
        raise ExcelKilitli from e
    ws = wb["Faturalar"]
    satir = ws.max_row + 1
    kayit = dict(kayit, kayit_no=satir - 1, kayit_tarihi=datetime.date.today())
    for anahtar, _b, _g, fmt in SUTUNLAR:
        c = ws.cell(row=satir, column=IDX[anahtar], value=kayit.get(anahtar))
        c.number_format = fmt
        if anahtar == "arsiv_yolu":
            _linkle(c)
    ws.auto_filter.ref = f"A1:{SON_SUTUN}{satir}"
    try:
        wb.save(EXCEL_YOLU)
    except PermissionError as e:
        raise ExcelKilitli from e
    return satir - 1


if __name__ == "__main__":
    if "--linkle" in sys.argv:
        try:
            print(f"{eski_kayitlari_linkle()} kayıt tıklanabilir yapıldı.")
        except PermissionError:
            print("Excel dosyası açık. Kapatıp komutu tekrar çalıştırın.")
    elif EXCEL_YOLU.exists():
        print(f"{EXCEL_YOLU.name} zaten var, üzerine yazılmadı.")
    else:
        sablon_olustur()
        print(f"Şablon oluşturuldu: {EXCEL_YOLU}")
